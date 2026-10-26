"""考试题库 Excel 模板与批量导入服务。"""
from __future__ import annotations

import json
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import AuthContext
from app.core.response import BizError
from app.models.exam import Question


HEADERS = [
    "题型",
    "题干",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
    "选项F",
    "选项G",
    "选项H",
    "正确答案",
    "解析",
    "分值",
    "难度",
    "分类ID",
    "标签",
]

ALLOWED_TYPES = {"single_choice", "multi_choice", "true_false", "essay"}
ALLOWED_DIFFICULTY = {"easy", "medium", "hard"}


def build_question_template() -> bytes:
    """生成 Excel 导入模板。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "题库导入"
    sheet.append(HEADERS)

    examples = [
        [
            "single_choice",
            "根据《安全生产法》，从业人员有权对本单位安全生产工作中存在的问题提出批评、检举和控告。",
            "正确",
            "错误",
            "",
            "",
            "",
            "",
            "",
            "",
            "A",
            "解析内容",
            2,
            "easy",
            1,
            "法律法规,从业人员",
        ],
        [
            "multi_choice",
            "下列属于高处作业防护用品的有哪些？",
            "安全带",
            "安全帽",
            "安全网",
            "普通布鞋",
            "",
            "",
            "",
            "",
            "ABC",
            "解析内容",
            3,
            "medium",
            2,
            "高处作业",
        ],
        [
            "true_false",
            "高处作业时，安全带应高挂低用。",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "正确",
            "解析内容",
            2,
            "easy",
            2,
            "高处作业",
        ],
    ]
    for row in examples:
        sheet.append(row)

    sheet.auto_filter.ref = f"A1:P{sheet.max_row}"

    type_validation = DataValidation(
        type="list",
        formula1='"single_choice,multi_choice,true_false,essay"',
        allow_blank=False,
    )
    type_validation.add("A2:A1000")
    sheet.add_data_validation(type_validation)

    difficulty_validation = DataValidation(
        type="list",
        formula1='"easy,medium,hard"',
        allow_blank=False,
    )
    difficulty_validation.add("N2:N1000")
    sheet.add_data_validation(difficulty_validation)

    sheet.freeze_panes = "A2"
    widths = [16, 50, 12, 12, 12, 12, 12, 12, 12, 12, 14, 30, 10, 12, 10, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell(row: tuple, index: int) -> str:
    if index >= len(row):
        return ""
    value = row[index]
    return "" if value is None else str(value).strip()


def _split_tags(value: str) -> list[str]:
    if not value:
        return []
    normalized = value.replace("，", ",").replace("；", ",").replace(";", ",")
    return [item.strip() for item in normalized.split(",") if item.strip()]


def _validate_row(row: tuple, index: int) -> tuple[dict | None, str | None]:
    q_type = _cell(row, 0)
    content = _cell(row, 1)
    answer = _cell(row, 10)
    explanation = _cell(row, 11)
    difficulty = _cell(row, 13) or "easy"
    category_raw = _cell(row, 14)
    tags = _split_tags(_cell(row, 15))

    if not content:
        return None, "题干不能为空"
    if q_type not in ALLOWED_TYPES:
        return None, f"题型'{q_type}'不支持"
    if len(content) > 2000:
        return None, "题干长度超过2000字限制"

    options = {}
    for option_index, key in enumerate("ABCDEFGH", start=2):
        value = _cell(row, option_index)
        if value:
            options[key] = value

    if q_type in {"single_choice", "multi_choice"}:
        if len(options) < 2:
            return None, "客观题至少需要两个选项"
        if not answer or any(key not in options for key in answer):
            return None, f"正确答案'{answer}'不在选项范围内"
    elif q_type == "true_false":
        if answer not in {"正确", "错误"}:
            return None, "判断题正确答案必须为正确或错误"
    elif q_type == "essay" and not answer:
        return None, "简答题必须填写参考答案"

    try:
        score = float(_cell(row, 12) or 2)
        if score <= 0:
            raise ValueError
    except ValueError:
        return None, "分值必须为正数"

    if difficulty not in ALLOWED_DIFFICULTY:
        return None, f"难度'{difficulty}'不支持"

    category_id = None
    if category_raw:
        try:
            category_id = int(category_raw)
        except ValueError:
            return None, "分类ID必须为整数"

    return (
        {
            "type": q_type,
            "content": content,
            "options": options or None,
            "answer": answer,
            "explanation": explanation,
            "score": score,
            "difficulty": difficulty,
            "category_id": category_id,
            "tags": tags,
        },
        None,
    )


def _normalize_row(header_row: tuple, row: tuple) -> tuple:
    normalized = [""] * 16
    option_map = {f"选项{key}": index for index, key in enumerate("ABCDEFGH", start=2)}
    for index, header in enumerate(header_row):
        name = _cell(header_row, index)
        if name == "题型":
            normalized[0] = _cell(row, index)
        elif name == "题干":
            normalized[1] = _cell(row, index)
        elif name in option_map:
            normalized[option_map[name]] = _cell(row, index)
        elif name == "正确答案":
            normalized[10] = _cell(row, index)
        elif name == "解析":
            normalized[11] = _cell(row, index)
        elif name == "分值":
            normalized[12] = _cell(row, index)
        elif name == "难度":
            normalized[13] = _cell(row, index)
        elif name in {"分类ID", "分类"}:
            normalized[14] = _cell(row, index)
        elif name == "标签":
            normalized[15] = _cell(row, index)
    return tuple(normalized)


async def import_questions_from_excel(
    db: AsyncSession,
    ctx: AuthContext,
    file_bytes: bytes,
    enter_bank: bool = True,
) -> dict:
    """解析 Excel 文件，正确行入库，错误行返回逐行原因。"""
    try:
        workbook = load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception:
        raise BizError(40001, "Excel 文件无法解析", 400)

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise BizError(40001, "Excel 文件为空", 400)

    header_row = rows[0]

    total = 0
    success = 0
    errors: list[dict] = []
    questions: list[Question] = []
    source_value = "import" if enter_bank else "paper_attachment"

    for row_index, row in enumerate(rows[1:], start=2):
        if not row or not any(cell is not None and str(cell).strip() for cell in row):
            continue
        total += 1
        normalized = _normalize_row(header_row, row)
        data, reason = _validate_row(normalized, row_index)
        if reason:
            errors.append({"row": row_index, "reason": reason})
            continue

        questions.append(
            Question(
                type=data["type"],
                content=data["content"],
                options_json=json.dumps(data["options"], ensure_ascii=False)
                if data["options"]
                else None,
                answer=data["answer"],
                explanation=data["explanation"],
                score=data["score"],
                difficulty=data["difficulty"],
                category_id=data["category_id"],
                tags_json=json.dumps(data["tags"], ensure_ascii=False),
                status="published" if ctx.role == "Admin" else "pending_review",
                source=source_value,
                created_by=ctx.user_id,
            )
        )
        success += 1

    question_ids: list[int] = []
    if questions:
        db.add_all(questions)
        await db.commit()
        question_ids = [q.question_id for q in questions]

    return {
        "total": total,
        "success": success,
        "failed": len(errors),
        "errors": errors,
        "question_ids": question_ids,
    }
